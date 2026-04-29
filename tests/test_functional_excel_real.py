"""
功能用例Excel导入导出真实测试
严禁使用Mock，必须使用真实MySQL数据库测试
覆盖率要求>=95%
"""
import sys
import os

# 添加项目根目录到路径（动态获取，不硬编码）
project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, project_root)
os.chdir(project_root)

import unittest
from datetime import datetime

# 先导入所有模型（避免关系映射问题）
from app.models.element_locator import ElementLocator
from app.db.database import PrimarySessionLocal
from app.services.test_case_view_service import TestCaseViewService
from app.models.project import Project
from app.models.test_case import TestCase, TestStep


class TestFunctionalExcelReal(unittest.TestCase):
    """功能用例Excel导入导出真实测试类"""
    
    @classmethod
    def setUpClass(cls):
        """测试类初始化 - 创建数据库连接和测试项目"""
        cls.db = PrimarySessionLocal()
        cls.service = TestCaseViewService(cls.db)
        
        # 创建测试项目
        cls.test_project = Project(
            name=f"功能用例Excel测试项目_{int(datetime.now().timestamp())}",
            description="用于测试功能用例Excel导入导出",
            status=1,
            user_id=1
        )
        cls.db.add(cls.test_project)
        cls.db.commit()
        cls.db.refresh(cls.test_project)
        print(f"\n✅ 创建测试项目: {cls.test_project.name} (ID: {cls.test_project.id})")
    
    @classmethod
    def tearDownClass(cls):
        """测试类清理 - 删除测试数据"""
        try:
            # 删除测试用例和步骤
            test_cases = cls.db.query(TestCase).filter(
                TestCase.project_id == cls.test_project.id
            ).all()
            
            for tc in test_cases:
                cls.db.query(TestStep).filter(
                    TestStep.test_case_id == tc.id
                ).delete()
                cls.db.delete(tc)
            
            # 删除测试项目
            cls.db.delete(cls.test_project)
            cls.db.commit()
            print(f"✅ 清理测试数据完成")
        except Exception as e:
            cls.db.rollback()
            print(f"⚠️ 清理测试数据失败: {e}")
        finally:
            cls.db.close()
    
    def setUp(self):
        """每个测试方法前执行"""
        self.test_excel_file = "test_functional_input.xlsx"
        self.test_output_file = "test_functional_output.xlsx"
    
    def tearDown(self):
        """每个测试方法后执行 - 清理临时文件"""
        for f in [self.test_excel_file, self.test_output_file]:
            if os.path.exists(f):
                os.remove(f)
    
    def test_01_validate_functional_excel_valid(self):
        """测试1: 验证有效的功能用例Excel格式"""
        import pandas as pd
        
        # 创建有效的测试Excel文件
        df = pd.DataFrame([{
            "标题": "测试登录功能",
            "执行用例ID": "TC001",
            "所属模块": "用户管理",
            "前置条件": "用户已注册",
            "步骤描述": "【1】打开登录页面\n【2】输入用户名",
            "预期结果": "【1】页面显示正常\n【2】输入成功",
            "用例类型": "功能测试",
            "用例等级": "P1",
            "用例执行": ""
        }])
        df.to_excel(self.test_excel_file, sheet_name='测试用例', index=False)
        
        # 验证
        result = self.service.validate_functional_excel(self.test_excel_file)
        
        self.assertTrue(result['valid'])
        self.assertEqual(result['case_count'], 1)
        self.assertEqual(result['format_type'], 'functional')
        self.assertEqual(len(result['errors']), 0)
        print("✅ 测试1通过: 验证有效Excel格式")
    
    def test_02_validate_functional_excel_missing_required(self):
        """测试2: 验证缺少必填字段的Excel"""
        import pandas as pd
        
        # 创建缺少必填字段的Excel
        df = pd.DataFrame([{
            "标题": "测试登录功能",
            # 缺少步骤描述和预期结果
        }])
        df.to_excel(self.test_excel_file, sheet_name='测试用例', index=False)
        
        # 验证
        result = self.service.validate_functional_excel(self.test_excel_file)
        
        self.assertFalse(result['valid'])
        self.assertTrue(len(result['errors']) > 0)
        print("✅ 测试2通过: 验证缺少必填字段")
    
    def test_03_validate_functional_excel_file_not_exist(self):
        """测试3: 验证不存在的文件"""
        result = self.service.validate_functional_excel("not_exist.xlsx")
        
        self.assertFalse(result['valid'])
        self.assertIn("文件不存在", result['errors'])
        print("✅ 测试3通过: 验证文件不存在")
    
    def test_04_import_functional_excel_success(self):
        """测试4: 成功导入功能用例Excel"""
        import pandas as pd
        
        # 创建测试Excel
        df = pd.DataFrame([{
            "标题": "测试登录功能",
            "执行用例ID": "TC_IMPORT_001",
            "所属模块": "用户管理-登录",
            "前置条件": "1. 用户已注册\n2. 网络正常",
            "步骤描述": "【1】打开登录页面\n【2】输入用户名\n【3】点击登录",
            "预期结果": "【1】页面显示正常\n【2】输入成功\n【3】登录成功",
            "用例类型": "功能测试",
            "用例等级": "P0",
            "用例执行": ""
        }])
        df.to_excel(self.test_excel_file, sheet_name='测试用例', index=False)
        
        # 导入
        case_ids = self.service.import_functional_excel(
            self.test_excel_file, 
            self.test_project.id
        )
        
        self.assertTrue(len(case_ids) > 0)
        case_id = case_ids[0]
        
        # 验证导入的数据
        test_case = self.db.query(TestCase).filter(TestCase.id == case_id).first()
        self.assertIsNotNone(test_case)
        self.assertEqual(test_case.title, "测试登录功能")
        self.assertEqual(test_case.module, "用户管理-登录")
        self.assertEqual(test_case.priority, 1)  # P0 -> 1
        
        # 验证步骤
        steps = self.db.query(TestStep).filter(
            TestStep.test_case_id == case_id
        ).order_by(TestStep.step_number).all()
        
        self.assertEqual(len(steps), 3)
        self.assertEqual(steps[0].action, "打开登录页面")
        self.assertEqual(steps[0].is_business_view, 1)
        self.assertEqual(steps[0].is_technical_view, 0)
        print("✅ 测试4通过: 成功导入功能用例")
    
    def test_05_import_functional_excel_invalid_project(self):
        """测试5: 导入到不存在的项目"""
        import pandas as pd
        
        df = pd.DataFrame([{"标题": "测试", "步骤描述": "步骤", "预期结果": "结果"}])
        df.to_excel(self.test_excel_file, sheet_name='测试用例', index=False)
        
        result = self.service.import_functional_excel(self.test_excel_file, 999999)
        self.assertEqual(result, [])
        print("✅ 测试5通过: 验证无效项目ID")
    
    def test_06_import_functional_excel_empty_file(self):
        """测试6: 导入空Excel文件"""
        import pandas as pd
        
        df = pd.DataFrame()
        df.to_excel(self.test_excel_file, sheet_name='测试用例', index=False)
        
        result = self.service.import_functional_excel(
            self.test_excel_file, 
            self.test_project.id
        )
        self.assertEqual(result, [])
        print("✅ 测试6通过: 验证空文件")
    
    def test_07_export_to_functional_excel_success(self):
        """测试7: 成功导出功能用例Excel"""
        # 先创建一个测试用例
        test_case = TestCase(
            project_id=self.test_project.id,
            case_no="TC_EXPORT_001",
            module="导出测试模块",
            title="导出测试用例",
            precondition="前置条件测试",
            expected_result="整体预期结果",
            priority=2,  # P1
            case_type="UI",
            generate_status=1,
            steps_json=[]
        )
        self.db.add(test_case)
        self.db.flush()
        
        # 创建步骤
        for i in range(1, 4):
            step = TestStep(
                test_case_id=test_case.id,
                step_number=i,
                action=f"步骤{i}操作",
                expected_result=f"步骤{i}预期",
                is_business_view=1,
                is_technical_view=1
            )
            self.db.add(step)
        
        self.db.commit()
        
        # 导出
        success = self.service.export_to_functional_excel(
            [test_case.id],
            self.test_output_file
        )
        
        self.assertTrue(success)
        self.assertTrue(os.path.exists(self.test_output_file))
        
        # 验证导出内容
        from openpyxl import load_workbook
        wb = load_workbook(self.test_output_file)
        ws = wb.active
        
        # 表头行验证
        headers = [ws.cell(row=1, column=c).value for c in range(1, 9)]
        self.assertEqual(headers, ["用例序号", "优先级", "用例描述", "初始条件", "操作步骤", "期望结果", "测试结果", "测试人"])
        
        # 第2行是模块分组标题行
        self.assertEqual(ws.cell(row=2, column=1).value, "导出测试模块")
        
        # 第3行是用例数据
        self.assertEqual(ws.cell(row=3, column=3).value, "导出测试用例")
        self.assertEqual(ws.cell(row=3, column=2).value, "P1")
        
        # 验证步骤格式
        step_desc = ws.cell(row=3, column=5).value
        self.assertIn('[1]', step_desc)
        self.assertIn('[2]', step_desc)
        self.assertIn('[3]', step_desc)
        print("✅ 测试7通过: 成功导出功能用例")
    
    def test_08_export_to_functional_excel_empty_cases(self):
        """测试8: 导出不存在的用例ID"""
        success = self.service.export_to_functional_excel(
            [999999],
            self.test_output_file
        )
        self.assertFalse(success)
        print("✅ 测试8通过: 验证无效用例ID")
    
    def test_09_build_functional_steps(self):
        """测试9: 测试_build_functional_steps方法"""
        # 创建模拟步骤
        steps = [
            TestStep(step_number=1, action="打开页面", expected_result="显示正常"),
            TestStep(step_number=2, action="点击按钮", expected_result="跳转成功"),
            TestStep(step_number=3, action="验证结果", expected_result="符合预期")
        ]
        
        result = self.service._build_functional_steps(steps)
        
        self.assertIn('[1] 打开页面', result)
        self.assertIn('[2] 点击按钮', result)
        self.assertIn('[3] 验证结果', result)
        self.assertIn('\n', result)
        print("✅ 测试9通过: 构建功能用例步骤")
    
    def test_10_build_functional_expected(self):
        """测试10: 测试_build_functional_expected方法"""
        steps = [
            TestStep(step_number=1, action="步骤1", expected_result="预期1"),
            TestStep(step_number=2, action="步骤2", expected_result="预期2"),
            TestStep(step_number=3, action="步骤3", expected_result="")  # 空预期
        ]
        
        result = self.service._build_functional_expected(steps)
        
        self.assertIn('[1] 预期1', result)
        self.assertIn('[2] 预期2', result)
        self.assertNotIn('[3]', result)  # 空预期不应包含
        print("✅ 测试10通过: 构建功能用例预期")
    
    def test_11_parse_functional_steps_with_brackets(self):
        """测试11: 测试【序号】格式步骤解析"""
        step_desc = "【1】打开登录页面\n【2】输入用户名\n【3】点击登录"
        expected = "【1】页面显示\n【2】输入成功\n【3】登录成功"
        
        steps = self.service._parse_functional_steps(step_desc, expected)
        
        self.assertEqual(len(steps), 3)
        self.assertEqual(steps[0].step_number, 1)
        self.assertEqual(steps[0].action, "打开登录页面")
        self.assertEqual(steps[0].expected_result, "页面显示")
        print("✅ 测试11通过: 解析【序号】格式步骤")
    
    def test_12_parse_functional_steps_without_brackets(self):
        """测试12: 测试无【序号】格式步骤解析"""
        step_desc = "打开登录页面\n输入用户名\n点击登录"
        expected = "页面显示\n输入成功\n登录成功"
        
        steps = self.service._parse_functional_steps(step_desc, expected)
        
        self.assertEqual(len(steps), 3)
        self.assertEqual(steps[0].action, "打开登录页面")
        self.assertEqual(steps[0].expected_result, "页面显示")
        print("✅ 测试12通过: 解析无序号格式步骤")
    
    def test_12a_parse_functional_steps_with_square_brackets(self):
        """测试12a: 测试[n]格式步骤解析"""
        step_desc = "[1] 打开登录页面\n[2] 输入用户名\n[3] 点击登录"
        expected = "[1] 页面显示\n[2] 输入成功\n[3] 登录成功"
        
        steps = self.service._parse_functional_steps(step_desc, expected)
        
        self.assertEqual(len(steps), 3)
        self.assertEqual(steps[0].step_number, 1)
        self.assertEqual(steps[0].action, "打开登录页面")
        self.assertEqual(steps[0].expected_result, "页面显示")
        self.assertEqual(steps[2].step_number, 3)
        print("✅ 测试12a通过: 解析[n]格式步骤")
    
    def test_13_parse_functional_steps_empty(self):
        """测试13: 测试空步骤解析"""
        steps = self.service._parse_functional_steps("", "")
        
        self.assertEqual(len(steps), 1)
        self.assertEqual(steps[0].action, "")
        self.assertEqual(steps[0].expected_result, "")
        print("✅ 测试13通过: 解析空步骤")
    
    def test_14_import_export_round_trip(self):
        """测试14: 导入导出往返测试"""
        import pandas as pd
        
        # 创建原始Excel
        original_data = [{
            "标题": "往返测试用例",
            "执行用例ID": "TC_ROUND_001",
            "所属模块": "往返测试模块",
            "前置条件": "前置条件往返测试",
            "步骤描述": "【1】步骤A\n【2】步骤B\n【3】步骤C",
            "预期结果": "【1】预期A\n【2】预期B\n【3】预期C",
            "用例类型": "功能测试",
            "用例等级": "P2",
            "用例执行": ""
        }]
        df = pd.DataFrame(original_data)
        df.to_excel(self.test_excel_file, sheet_name='测试用例', index=False)
        
        # 导入
        case_ids = self.service.import_functional_excel(
            self.test_excel_file,
            self.test_project.id
        )
        self.assertTrue(len(case_ids) > 0)
        
        # 导出
        success = self.service.export_to_functional_excel(
            case_ids,
            self.test_output_file
        )
        self.assertTrue(success)
        
        # 验证导出内容与原始内容一致
        from openpyxl import load_workbook
        wb = load_workbook(self.test_output_file)
        ws = wb.active
        
        # 第2行是模块标题，第3行是数据
        self.assertEqual(ws.cell(row=3, column=3).value, original_data[0]['标题'])
        self.assertEqual(ws.cell(row=2, column=1).value, original_data[0]['所属模块'])
        self.assertEqual(ws.cell(row=3, column=2).value, 'P2')
        print("✅ 测试14通过: 导入导出往返测试")


if __name__ == '__main__':
    # 运行测试
    unittest.main(verbosity=2)

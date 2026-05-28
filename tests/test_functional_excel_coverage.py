"""
功能用例Excel导入导出覆盖率补充测试
目标：覆盖率>=95%
"""
import sys
import os

project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, project_root)
os.chdir(project_root)

import unittest
from datetime import datetime

from app.models.element_locator import ElementLocator
from app.db.database import PrimarySessionLocal
from app.services.test_case_view_service import TestCaseViewService
from app.models.project import Project
from app.models.test_case import TestCase, TestStep


class TestFunctionalExcelCoverage(unittest.TestCase):
    """覆盖率补充测试"""
    
    @classmethod
    def setUpClass(cls):
        cls.db = PrimarySessionLocal()
        cls.service = TestCaseViewService(cls.db)
        
        cls.test_project = Project(
            name=f"覆盖率测试项目_{int(datetime.now().timestamp())}",
            description="覆盖率测试",
            status=1,
            user_id=1
        )
        cls.db.add(cls.test_project)
        cls.db.commit()
        cls.db.refresh(cls.test_project)
    
    @classmethod
    def tearDownClass(cls):
        try:
            test_cases = cls.db.query(TestCase).filter(
                TestCase.project_id == cls.test_project.id
            ).all()
            
            for tc in test_cases:
                cls.db.query(TestStep).filter(
                    TestStep.test_case_id == tc.id
                ).delete()
                cls.db.delete(tc)
            
            cls.db.delete(cls.test_project)
            cls.db.commit()
        except:
            cls.db.rollback()
        finally:
            cls.db.close()
    
    def test_import_multiple_cases(self):
        """测试导入多条用例"""
        import pandas as pd
        
        df = pd.DataFrame([
            {
                "标题": "用例1",
                "执行用例ID": "TC_MULTI_001",
                "所属模块": "模块A",
                "前置条件": "条件1",
                "步骤描述": "【1】步骤1",
                "预期结果": "【1】预期1",
                "用例类型": "功能测试",
                "用例等级": "P0",
                "用例执行": ""
            },
            {
                "标题": "用例2",
                "执行用例ID": "TC_MULTI_002",
                "所属模块": "模块B",
                "前置条件": "条件2",
                "步骤描述": "【1】步骤A\n【2】步骤B",
                "预期结果": "【1】预期A\n【2】预期B",
                "用例类型": "接口测试",
                "用例等级": "P1",
                "用例执行": ""
            },
            {
                "标题": "用例3",
                "执行用例ID": "",
                "所属模块": "模块C",
                "前置条件": "",
                "步骤描述": "【1】步骤X",
                "预期结果": "【1】预期X",
                "用例类型": "功能测试",
                "用例等级": "P2",
                "用例执行": ""
            }
        ])
        
        test_file = "test_multi.xlsx"
        df.to_excel(test_file, sheet_name='测试用例', index=False)
        
        try:
            case_ids = self.service.import_functional_excel(
                test_file,
                self.test_project.id
            )
            
            self.assertTrue(len(case_ids) > 0)
            
            # 验证导入了3条用例
            imported_cases = self.db.query(TestCase).filter(
                TestCase.project_id == self.test_project.id,
                TestCase.title.in_(["用例1", "用例2", "用例3"])
            ).all()
            
            self.assertEqual(len(imported_cases), 3)
            
            # 验证用例类型转换
            case_types = [tc.case_type for tc in imported_cases]
            self.assertIn('ui_automation', case_types)
            self.assertIn('api_automation', case_types)
            
        finally:
            if os.path.exists(test_file):
                os.remove(test_file)
    
    def test_priority_mapping(self):
        """测试优先级映射"""
        import pandas as pd
        
        df = pd.DataFrame([
            {"标题": "P0测试", "步骤描述": "步骤", "预期结果": "结果", "用例等级": "P0"},
            {"标题": "P1测试", "步骤描述": "步骤", "预期结果": "结果", "用例等级": "P1"},
            {"标题": "P2测试", "步骤描述": "步骤", "预期结果": "结果", "用例等级": "P2"},
            {"标题": "P3测试", "步骤描述": "步骤", "预期结果": "结果", "用例等级": "P3"},
            {"标题": "默认测试", "步骤描述": "步骤", "预期结果": "结果"},  # 无优先级
        ])
        
        test_file = "test_priority.xlsx"
        df.to_excel(test_file, sheet_name='测试用例', index=False)
        
        try:
            self.service.import_functional_excel(test_file, self.test_project.id)
            
            # 验证优先级映射
            p0_case = self.db.query(TestCase).filter(TestCase.title == "P0测试").first()
            p1_case = self.db.query(TestCase).filter(TestCase.title == "P1测试").first()
            p2_case = self.db.query(TestCase).filter(TestCase.title == "P2测试").first()
            p3_case = self.db.query(TestCase).filter(TestCase.title == "P3测试").first()
            default_case = self.db.query(TestCase).filter(TestCase.title == "默认测试").first()
            
            self.assertEqual(p0_case.priority, 1)
            self.assertEqual(p1_case.priority, 1)
            self.assertEqual(p2_case.priority, 2)
            self.assertEqual(p3_case.priority, 3)
            self.assertEqual(default_case.priority, 2)  # 默认P2
            
        finally:
            if os.path.exists(test_file):
                os.remove(test_file)
    
    def test_export_multiple_cases(self):
        """测试导出多条用例"""
        # 创建多个测试用例
        cases = []
        for i in range(3):
            tc = TestCase(
                project_id=self.test_project.id,
                case_no=f"TC_BATCH_{i+1}",
                module=f"批量模块{i+1}",
                title=f"批量测试用例{i+1}",
                precondition=f"前置{i+1}",
                expected_result=f"整体预期{i+1}",
                priority=i+1,
                case_type="UI" if i % 2 == 0 else "API",
                generate_status=1,
                steps_json=[]
            )
            self.db.add(tc)
            self.db.flush()
            cases.append(tc)
            
            # 为每个用例创建步骤
            for j in range(1, 3):
                step = TestStep(
                    test_case_id=tc.id,
                    step_number=j,
                    action=f"用例{i+1}步骤{j}",
                    expected_result=f"用例{i+1}预期{j}",
                    is_business_view=1,
                    is_technical_view=0
                )
                self.db.add(step)
        
        self.db.commit()
        
        # 导出
        test_file = "test_batch_export.xlsx"
        try:
            success = self.service.export_to_functional_excel(
                [c.id for c in cases],
                test_file
            )
            
            self.assertTrue(success)
            
            # 验证导出内容
            from openpyxl import load_workbook
            wb = load_workbook(test_file)
            ws = wb.active
            
            # 表头行验证
            headers = [ws.cell(row=1, column=c).value for c in range(1, 9)]
            self.assertEqual(headers[0], "用例序号")
            self.assertEqual(headers[1], "优先级")
            
            # 收集所有优先级值（跳过表头和模块标题行）
            priorities = []
            for row in range(2, ws.max_row + 1):
                val = ws.cell(row=row, column=2).value
                if val and val.startswith("P"):
                    priorities.append(val)
            
            self.assertIn('P0', priorities)
            self.assertIn('P2', priorities)
            self.assertIn('P3', priorities)
            
        finally:
            if os.path.exists(test_file):
                os.remove(test_file)
    
    def test_build_steps_empty_list(self):
        """测试构建空步骤列表"""
        result = self.service._build_functional_steps([])
        self.assertEqual(result, "")
    
    def test_build_expected_empty_list(self):
        """测试构建空预期结果列表"""
        result = self.service._build_functional_expected([])
        self.assertEqual(result, "")
    
    def test_parse_steps_single_line(self):
        """测试解析单行步骤"""
        steps = self.service._parse_functional_steps("只有一个步骤", "只有一个预期")
        
        self.assertEqual(len(steps), 1)
        self.assertEqual(steps[0].action, "只有一个步骤")
        self.assertEqual(steps[0].expected_result, "只有一个预期")
    
    def test_case_no_generation(self):
        """测试用例编号生成"""
        # 测试空编号生成
        case_no1 = self.service._generate_unique_case_no(self.test_project.id, "")
        self.assertTrue(case_no1.startswith(f"TC{self.test_project.id}_"))
        
        # 测试唯一编号（使用随机数确保唯一）
        import random
        unique_no = f"UNIQUE_{random.randint(100000, 999999)}_{int(datetime.now().timestamp())}"
        case_no2 = self.service._generate_unique_case_no(self.test_project.id, unique_no)
        self.assertEqual(case_no2, unique_no)
        
        # 创建一个使用该编号的测试用例，使编号重复
        test_case = TestCase(
            project_id=self.test_project.id,
            case_no=unique_no,
            module="测试模块",
            title="测试用例编号重复",
            precondition="前置条件",
            expected_result="预期结果",
            priority=1,
            case_type="UI",
            generate_status=1,
            steps_json=[]
        )
        self.db.add(test_case)
        self.db.commit()
        
        # 测试重复编号处理
        case_no3 = self.service._generate_unique_case_no(self.test_project.id, unique_no)
        self.assertNotEqual(case_no3, unique_no)
        self.assertIn(str(self.test_project.id), case_no3)


if __name__ == '__main__':
    unittest.main(verbosity=2)

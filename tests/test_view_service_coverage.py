"""
TestCaseViewService 覆盖率补充测试
目标：将覆盖率从48%提升到95%以上
严禁使用Mock，必须使用真实MySQL数据库
"""
import sys
import os
import unittest
from datetime import datetime
import pytest

pytestmark = pytest.mark.skip(reason="数据库DDL不兼容")

project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, project_root)
os.chdir(project_root)

from app.db.database import PrimarySessionLocal
from app.services.test_case_view_service import TestCaseViewService
from app.models.project import Project
from app.models.test_case import TestCase, TestStep
from app.models.element_locator import ElementLocator


class TestViewServiceCoverage(unittest.TestCase):
    """补充测试用例，提高覆盖率"""
    
    @classmethod
    def setUpClass(cls):
        """测试类初始化"""
        cls.db = PrimarySessionLocal()
        cls.service = TestCaseViewService(cls.db)
        
        # 创建测试项目
        cls.test_project = Project(
            name=f"覆盖率测试项目_{int(datetime.now().timestamp())}",
            description="用于覆盖率测试",
            status=1,
            user_id=1
        )
        cls.db.add(cls.test_project)
        cls.db.commit()
        cls.db.refresh(cls.test_project)
        print(f"\n✅ 创建测试项目: {cls.test_project.name} (ID: {cls.test_project.id})")
        
        # 创建测试用例
        cls.test_case = TestCase(
            project_id=cls.test_project.id,
            case_no=f"TC_COV_{int(datetime.now().timestamp())}",
            module="覆盖率测试模块",
            title="覆盖率测试用例",
            precondition="前置条件",
            expected_result="预期结果",
            priority=1,
            case_type="UI",
            generate_status=1,
            steps_json=[]
        )
        cls.db.add(cls.test_case)
        cls.db.flush()
        
        # 创建测试步骤
        for i in range(1, 4):
            step = TestStep(
                test_case_id=cls.test_case.id,
                step_number=i,
                action=f"操作步骤{i}",
                expected_result=f"预期结果{i}",
                is_business_view=1,
                is_technical_view=1,
                has_locator=1 if i == 1 else 0,
                locator_status="located" if i == 1 else "pending"
            )
            cls.db.add(step)
            cls.db.flush()
            
            # 为第一个步骤添加定位器
            if i == 1:
                locator = ElementLocator(
                    step_id=step.id,
                    css_selector=f"#element{i}",
                    xpath=f"//element[{i}]",
                    element_type="button",
                    ai_confidence=0.95
                )
                cls.db.add(locator)
        
        cls.db.commit()
        print(f"✅ 创建测试用例: {cls.test_case.title} (ID: {cls.test_case.id})")
    
    @classmethod
    def tearDownClass(cls):
        """测试类清理"""
        # 清理测试数据
        cls.db.query(ElementLocator).filter(
            ElementLocator.step_id.in_(
                cls.db.query(TestStep.id).filter(TestStep.test_case_id == cls.test_case.id)
            )
        ).delete(synchronize_session=False)
        cls.db.query(TestStep).filter(TestStep.test_case_id == cls.test_case.id).delete(synchronize_session=False)
        cls.db.query(TestCase).filter(TestCase.id == cls.test_case.id).delete(synchronize_session=False)
        cls.db.query(Project).filter(Project.id == cls.test_project.id).delete(synchronize_session=False)
        cls.db.commit()
        cls.db.close()
        print("\n✅ 清理测试数据完成")
    
    def test_01_export_to_excel(self):
        """测试导出Excel功能"""
        import tempfile
        import pandas as pd
        
        # 创建临时目录和文件
        temp_dir = tempfile.mkdtemp()
        temp_path = os.path.join(temp_dir, 'test_export.xlsx')
        
        try:
            # 测试导出
            result = self.service.export_to_excel(self.test_case.id, temp_path)
            self.assertTrue(result, "导出Excel应该成功")
            self.assertTrue(os.path.exists(temp_path), "Excel文件应该存在")
            
            # 验证文件内容
            xl = pd.ExcelFile(temp_path)
            self.assertIn('用例信息', xl.sheet_names, "应该有用例信息sheet")
            self.assertIn('测试步骤', xl.sheet_names, "应该有测试步骤sheet")
            
            # 验证用例信息
            case_info = pd.read_excel(temp_path, sheet_name='用例信息')
            self.assertEqual(len(case_info), 1, "应该有一条用例信息")
            self.assertEqual(case_info.iloc[0]['用例标题'], self.test_case.title)
            
            # 验证测试步骤
            steps = pd.read_excel(temp_path, sheet_name='测试步骤')
            self.assertEqual(len(steps), 3, "应该有3个步骤")
            
            print("✅ 测试1通过: 导出Excel功能")
        finally:
            # Windows下需要手动关闭文件句柄
            import shutil
            try:
                shutil.rmtree(temp_dir)
            except:
                pass
    
    def test_02_export_to_functional_excel(self):
        """测试导出功能用例Excel"""
        import tempfile
        import pandas as pd
        
        temp_dir = tempfile.mkdtemp()
        temp_path = os.path.join(temp_dir, 'test_functional.xlsx')
        
        try:
            result = self.service.export_to_functional_excel([self.test_case.id], temp_path)
            self.assertTrue(result, "导出功能用例Excel应该成功")
            self.assertTrue(os.path.exists(temp_path), "文件应该存在")
            
            # 验证内容
            from openpyxl import load_workbook
            wb = load_workbook(temp_path)
            ws = wb.active
            headers = [ws.cell(row=1, column=c).value for c in range(1, 9)]
            self.assertIn('用例描述', headers, "应该有用例描述列")
            self.assertIn('操作步骤', headers, "应该有操作步骤列")
            # 第2行是模块标题，第3行是用例数据
            self.assertEqual(ws.cell(row=3, column=3).value, self.test_case.title)
            
            print("✅ 测试2通过: 导出功能用例Excel")
        finally:
            import shutil
            try:
                shutil.rmtree(temp_dir)
            except:
                pass
    
    def test_03_import_from_excel(self):
        """测试从Excel导入"""
        import tempfile
        import pandas as pd
        
        # 创建测试Excel文件
        temp_dir = tempfile.mkdtemp()
        temp_path = os.path.join(temp_dir, 'test_import.xlsx')
        
        try:
            # 创建Excel内容
            with pd.ExcelWriter(temp_path, engine='openpyxl') as writer:
                # 用例信息
                case_df = pd.DataFrame([{
                    '用例编号': f'IMPORT_{int(datetime.now().timestamp())}',
                    '用例标题': '导入测试用例',
                    '所属模块': '导入测试模块',
                    '前置条件': '前置条件测试',
                    '预期结果': '预期结果测试',
                    '优先级': 1,
                    '总步骤数': 2
                }])
                case_df.to_excel(writer, sheet_name='用例信息', index=False)
                
                # 测试步骤
                steps_df = pd.DataFrame([
                    {
                        '步骤编号': 1,
                        '操作步骤': '第一步操作',
                        '预期结果': '第一步预期',
                        '业务视图': '是',
                        '技术视图': '是',
                        '已定位': '是',
                        '定位状态': 'located',
                        'CSS选择器': '#step1',
                        'XPath': '//step[1]',
                        '元素类型': 'button'
                    },
                    {
                        '步骤编号': 2,
                        '操作步骤': '第二步操作',
                        '预期结果': '第二步预期',
                        '业务视图': '是',
                        '技术视图': '否',
                        '已定位': '否',
                        '定位状态': 'pending',
                        'CSS选择器': '',
                        'XPath': '',
                        '元素类型': ''
                    }
                ])
                steps_df.to_excel(writer, sheet_name='测试步骤', index=False)
            
            # 测试导入
            case_id = self.service.import_from_excel(temp_path, self.test_project.id)
            self.assertIsNotNone(case_id, "导入应该成功并返回用例ID")
            
            # 验证导入的数据
            imported_case = self.db.query(TestCase).filter(TestCase.id == case_id).first()
            self.assertIsNotNone(imported_case, "导入的用例应该存在")
            self.assertEqual(imported_case.title, '导入测试用例')
            self.assertEqual(len(imported_case.test_steps), 2, "应该有2个步骤")
            
            # 清理导入的数据
            self.db.query(TestStep).filter(TestStep.test_case_id == case_id).delete(synchronize_session=False)
            self.db.query(TestCase).filter(TestCase.id == case_id).delete(synchronize_session=False)
            self.db.commit()
            
            print("✅ 测试3通过: 从Excel导入")
        finally:
            import shutil
            try:
                shutil.rmtree(temp_dir)
            except:
                pass
    
    def test_04_validate_excel_format(self):
        """测试验证Excel格式"""
        import tempfile
        import pandas as pd
        
        # 测试有效格式
        temp_dir = tempfile.mkdtemp()
        temp_path = os.path.join(temp_dir, 'test_validate.xlsx')
        
        try:
            with pd.ExcelWriter(temp_path, engine='openpyxl') as writer:
                # 用例信息需要包含必需的列
                pd.DataFrame([{
                    '用例编号': 'TC001',
                    '用例标题': '测试用例标题'
                }]).to_excel(writer, sheet_name='用例信息', index=False)
                # 测试步骤需要包含必需的列
                pd.DataFrame([{
                    '步骤编号': 1,
                    '操作步骤': '操作',
                    '预期结果': '预期'
                }]).to_excel(writer, sheet_name='测试步骤', index=False)
            
            result = self.service.validate_excel_format(temp_path)
            self.assertTrue(result['valid'], f"有效格式应该验证通过，错误: {result.get('errors')}")
            
            print("✅ 测试4通过: 验证Excel格式")
        finally:
            import shutil
            try:
                shutil.rmtree(temp_dir)
            except:
                pass
    
    def test_05_validate_functional_excel(self):
        """测试验证功能用例Excel格式"""
        import tempfile
        import pandas as pd
        
        temp_dir = tempfile.mkdtemp()
        temp_path = os.path.join(temp_dir, 'test_validate_func.xlsx')
        
        try:
            # 创建有效的功能用例Excel
            df = pd.DataFrame([{
                '标题': '测试用例',
                '执行用例ID': 'TC001',
                '所属模块': '模块1',
                '前置条件': '条件1',
                '步骤描述': '【1】步骤1',
                '预期结果': '【1】预期1',
                '用例类型': '功能测试',
                '用例等级': 'P0'
            }])
            df.to_excel(temp_path, index=False)
            
            result = self.service.validate_functional_excel(temp_path)
            self.assertTrue(result['valid'], "有效格式应该验证通过")
            
            print("✅ 测试5通过: 验证功能用例Excel格式")
        finally:
            import shutil
            try:
                shutil.rmtree(temp_dir)
            except:
                pass
    
    def test_06_export_technical_view_to_json(self):
        """测试导出技术视图为JSON"""
        result = self.service.export_technical_view_to_json(self.test_case.id)
        self.assertIsInstance(result, dict, "结果应该是字典")
        self.assertIn('case_id', result, "应该包含case_id")
        self.assertIn('steps', result, "应该包含steps")
        
        print("✅ 测试6通过: 导出技术视图为JSON")
    
    def test_07_export_technical_view_to_python(self):
        """测试导出技术视图为Python脚本"""
        result = self.service.export_technical_view_to_python(self.test_case.id)
        self.assertIsInstance(result, str, "结果应该是字符串")
        self.assertIn('import pytest', result, "应该包含pytest导入")
        self.assertIn('async def test_', result, "应该包含测试函数")
        
        print("✅ 测试7通过: 导出技术视图为Python脚本")
    
    def test_08_update_step_view_config(self):
        """测试更新步骤视图配置"""
        # 获取一个步骤
        step = self.db.query(TestStep).filter(TestStep.test_case_id == self.test_case.id).first()
        self.assertIsNotNone(step, "应该存在测试步骤")
        
        # 更新配置
        result = self.service.update_step_view_config(
            step.id,
            is_business_view=0,
            is_technical_view=0
        )
        self.assertTrue(result, "更新应该成功")
        
        # 验证更新
        self.db.refresh(step)
        self.assertEqual(step.is_business_view, 0)
        self.assertEqual(step.is_technical_view, 0)
        
        # 恢复
        result = self.service.update_step_view_config(
            step.id,
            is_business_view=1,
            is_technical_view=1
        )
        self.assertTrue(result, "恢复应该成功")
        
        print("✅ 测试8通过: 更新步骤视图配置")
    
    def test_09_batch_update_view_flags(self):
        """测试批量更新视图标记"""
        # 批量更新业务视图配置
        count = self.service.batch_update_view_config(
            self.test_case.id,
            view_type='business',
            visible=False
        )
        self.assertGreater(count, 0, "应该更新至少一个步骤")
        
        # 验证
        steps = self.db.query(TestStep).filter(TestStep.test_case_id == self.test_case.id).all()
        for step in steps:
            self.assertEqual(step.is_business_view, 0)
        
        # 恢复
        self.service.batch_update_view_config(self.test_case.id, view_type='business', visible=True)
        
        print("✅ 测试9通过: 批量更新视图标记")
    
    def test_10_get_view_statistics(self):
        """测试获取视图统计"""
        stats = self.service.get_view_statistics(self.test_case.id)
        
        self.assertIn('total_steps', stats, "应该包含total_steps")
        self.assertIn('business_view_steps', stats, "应该包含business_view_steps")
        self.assertIn('technical_view_steps', stats, "应该包含technical_view_steps")
        self.assertIn('located_steps', stats, "应该包含located_steps")
        self.assertIn('locator_coverage', stats, "应该包含locator_coverage")
        
        print("✅ 测试10通过: 获取视图统计")
    
    def test_11_nonexistent_case_operations(self):
        """测试对不存在用例的操作"""
        nonexistent_id = 999999
        
        # 测试各种方法对不存在用例的处理
        result = self.service.export_to_excel(nonexistent_id, '/tmp/test.xlsx')
        self.assertFalse(result, "导出不存在的用例应该失败")
        
        result = self.service.export_technical_view_to_json(nonexistent_id)
        self.assertEqual(result, {}, "不存在的用例应该返回空字典")
        
        result = self.service.export_technical_view_to_python(nonexistent_id)
        self.assertEqual(result, "", "不存在的用例应该返回空字符串")
        
        print("✅ 测试11通过: 不存在用例的操作处理")
    
    def test_12_invalid_excel_format(self):
        """测试无效Excel格式验证"""
        import tempfile
        import pandas as pd
        
        # 测试缺少必要sheet
        temp_dir = tempfile.mkdtemp()
        temp_path = os.path.join(temp_dir, 'test_invalid.xlsx')
        
        try:
            # 创建缺少sheet的Excel
            pd.DataFrame({'col': [1]}).to_excel(temp_path, index=False)
            
            result = self.service.validate_excel_format(temp_path)
            self.assertFalse(result['valid'], "无效格式应该验证失败")
            self.assertTrue(len(result['errors']) > 0, "应该有错误信息")
            
            print("✅ 测试12通过: 无效Excel格式验证")
        finally:
            import shutil
            try:
                shutil.rmtree(temp_dir)
            except:
                pass


    def test_13_export_business_markdown(self):
        """测试导出业务视图为Markdown"""
        result = self.service.export_business_view_to_markdown(self.test_case.id)
        self.assertIsInstance(result, str, "结果应该是字符串")
        self.assertIn(self.test_case.title, result, "应该包含用例标题")
        self.assertIn('## 测试步骤', result, "应该包含测试步骤标题")
        
        print("✅ 测试13通过: 导出业务视图为Markdown")
    
    def test_14_export_business_html(self):
        """测试导出业务视图为HTML"""
        result = self.service.export_business_view_to_html(self.test_case.id)
        self.assertIsInstance(result, str, "结果应该是字符串")
        self.assertIn('<!DOCTYPE html>', result, "应该是HTML格式")
        self.assertIn(self.test_case.title, result, "应该包含用例标题")
        
        print("✅ 测试14通过: 导出业务视图为HTML")
    
    def test_15_get_business_view(self):
        """测试获取业务视图"""
        view = self.service.get_business_view(self.test_case.id)
        self.assertIsNotNone(view, "业务视图不应该为空")
        self.assertEqual(view.case_id, self.test_case.id, "case_id应该匹配")
        self.assertEqual(view.title, self.test_case.title, "标题应该匹配")
        self.assertTrue(len(view.steps) > 0, "应该有步骤")
        
        print("✅ 测试15通过: 获取业务视图")
    
    def test_16_import_functional_excel(self):
        """测试从功能用例Excel导入"""
        import tempfile
        import pandas as pd
        
        temp_dir = tempfile.mkdtemp()
        temp_path = os.path.join(temp_dir, 'test_import_func.xlsx')
        
        try:
            # 创建功能用例Excel
            df = pd.DataFrame([{
                '标题': '功能导入测试',
                '执行用例ID': 'TC_FUNC_001',
                '所属模块': '功能模块',
                '前置条件': '前置条件',
                '步骤描述': '【1】步骤1\n【2】步骤2',
                '预期结果': '【1】预期1\n【2】预期2',
                '用例类型': '功能测试',
                '用例等级': 'P1'
            }])
            df.to_excel(temp_path, index=False)
            
            # 测试导入
            case_ids = self.service.import_functional_excel(temp_path, self.test_project.id, module="功能模块")
            self.assertTrue(len(case_ids) > 0, "导入应该成功")
            
            # 验证导入的数据
            imported_case = self.db.query(TestCase).filter(TestCase.id == case_ids[0]).first()
            self.assertIsNotNone(imported_case, "导入的用例应该存在")
            self.assertEqual(imported_case.title, '功能导入测试')
            
            # 清理
            for cid in case_ids:
                self.db.query(TestStep).filter(TestStep.test_case_id == cid).delete(synchronize_session=False)
                self.db.query(TestCase).filter(TestCase.id == cid).delete(synchronize_session=False)
            self.db.commit()
            
            print("✅ 测试16通过: 从功能用例Excel导入")
        finally:
            import shutil
            try:
                shutil.rmtree(temp_dir)
            except:
                pass
    
    def test_17_generate_unique_case_no(self):
        """测试生成唯一用例编号"""
        # 测试正常情况
        case_no = self.service._generate_unique_case_no(self.test_project.id, "TC_UNIQUE_001")
        self.assertIsInstance(case_no, str, "应该返回字符串")
        self.assertTrue(len(case_no) > 0, "编号不应该为空")
        
        # 测试重复编号处理
        case_no2 = self.service._generate_unique_case_no(self.test_project.id, self.test_case.case_no)
        self.assertIsInstance(case_no2, str, "应该返回字符串")
        
        print("✅ 测试17通过: 生成唯一用例编号")
    
    def test_18_empty_case_business_view(self):
        """测试空用例的业务视图"""
        # 创建一个没有步骤的用例
        empty_case = TestCase(
            project_id=self.test_project.id,
            case_no=f"TC_EMPTY_{int(datetime.now().timestamp())}",
            module="空用例模块",
            title="空用例测试",
            precondition="",
            expected_result="",
            priority=2,
            case_type="UI",
            generate_status=1,
            steps_json=[]
        )
        self.db.add(empty_case)
        self.db.commit()
        self.db.refresh(empty_case)
        
        try:
            # 测试业务视图
            view = self.service.get_business_view(empty_case.id)
            self.assertIsNotNone(view, "空用例的业务视图不应该为空")
            self.assertEqual(len(view.steps), 0, "空用例应该没有步骤")
            
            # 测试技术视图
            tech_view = self.service.get_technical_view(empty_case.id)
            self.assertIsNotNone(tech_view, "空用例的技术视图不应该为空")
            self.assertEqual(len(tech_view['steps']), 0, "空用例应该没有技术步骤")
            
            print("✅ 测试18通过: 空用例的业务视图")
        finally:
            self.db.query(TestCase).filter(TestCase.id == empty_case.id).delete(synchronize_session=False)
            self.db.commit()
    
    def test_19_locator_coverage_empty(self):
        """测试空用例的定位覆盖率"""
        # 创建一个没有步骤的用例
        empty_case = TestCase(
            project_id=self.test_project.id,
            case_no=f"TC_COV_EMPTY_{int(datetime.now().timestamp())}",
            module="覆盖率模块",
            title="覆盖率空用例",
            precondition="",
            expected_result="",
            priority=2,
            case_type="UI",
            generate_status=1,
            steps_json=[]
        )
        self.db.add(empty_case)
        self.db.commit()
        self.db.refresh(empty_case)
        
        try:
            coverage = self.service.get_locator_coverage(empty_case.id)
            self.assertEqual(coverage['total_steps'], 0, "总步骤数应该为0")
            self.assertEqual(coverage['coverage_percentage'], 0.0, "覆盖率应该为0")
            
            print("✅ 测试19通过: 空用例的定位覆盖率")
        finally:
            self.db.query(TestCase).filter(TestCase.id == empty_case.id).delete(synchronize_session=False)
            self.db.commit()
    
    def test_20_business_vs_technical_view(self):
        """测试业务视图和技术视图的差异"""
        # 获取业务视图
        business_view = self.service.get_business_view(self.test_case.id)
        
        # 获取技术视图
        technical_view = self.service.get_technical_view(self.test_case.id)
        
        # 业务视图应该只包含 is_business_view=1 的步骤
        # 技术视图应该只包含 is_technical_view=1 的步骤
        self.assertIsNotNone(business_view)
        self.assertIsNotNone(technical_view)
        
        print("✅ 测试20通过: 业务视图和技术视图的差异")
    
    def test_21_technical_step_view_to_dict(self):
        """测试TechnicalStepView的to_dict方法"""
        from app.services.test_case_view_service import TechnicalStepView
        
        # 创建完整的技术步骤视图
        step_view = TechnicalStepView(
            step_number=1,
            action="点击按钮",
            expected_result="按钮被点击",
            has_locator=True,
            locator_status="located",
            css_selector="#button",
            xpath="//button",
            ai_coordinate={"x": 100, "y": 200},
            element_type="button"
        )
        
        result = step_view.to_dict()
        
        self.assertEqual(result['step_number'], 1)
        self.assertEqual(result['action'], "点击按钮")
        self.assertEqual(result['css_selector'], "#button")
        self.assertEqual(result['xpath'], "//button")
        self.assertEqual(result['ai_coordinate'], {"x": 100, "y": 200})
        self.assertEqual(result['element_type'], "button")
        
        print("✅ 测试21通过: TechnicalStepView.to_dict")
    
    def test_22_technical_step_view_to_dict_partial(self):
        """测试TechnicalStepView的to_dict方法（部分字段）"""
        from app.services.test_case_view_service import TechnicalStepView
        
        # 创建只有必填字段的技术步骤视图
        step_view = TechnicalStepView(
            step_number=2,
            action="输入文本",
            expected_result="文本已输入",
            has_locator=False,
            locator_status="pending"
        )
        
        result = step_view.to_dict()
        
        self.assertEqual(result['step_number'], 2)
        self.assertNotIn('css_selector', result)
        self.assertNotIn('xpath', result)
        self.assertNotIn('ai_coordinate', result)
        self.assertNotIn('element_type', result)
        
        print("✅ 测试22通过: TechnicalStepView.to_dict部分字段")
    
    def test_23_business_step_view_to_dict(self):
        """测试BusinessStepView的to_dict方法"""
        from app.services.test_case_view_service import BusinessStepView
        
        step_view = BusinessStepView(
            step_number=1,
            action="打开页面",
            expected_result="页面打开成功"
        )
        
        result = step_view.to_dict()
        
        self.assertEqual(result['step_number'], 1)
        self.assertEqual(result['action'], "打开页面")
        self.assertEqual(result['expected_result'], "页面打开成功")
        
        print("✅ 测试23通过: BusinessStepView.to_dict")
    
    def test_24_batch_update_technical_view(self):
        """测试批量更新技术视图配置"""
        # 先更新为技术视图不可见
        count = self.service.batch_update_view_config(
            self.test_case.id,
            view_type='technical',
            visible=False
        )
        self.assertGreater(count, 0, "应该更新至少一个步骤")
        
        # 验证
        steps = self.db.query(TestStep).filter(TestStep.test_case_id == self.test_case.id).all()
        for step in steps:
            self.assertEqual(step.is_technical_view, 0)
        
        # 恢复
        self.service.batch_update_view_config(self.test_case.id, view_type='technical', visible=True)
        
        print("✅ 测试24通过: 批量更新技术视图配置")
    
    def test_25_batch_update_invalid_view_type(self):
        """测试批量更新无效的视图类型"""
        count = self.service.batch_update_view_config(
            self.test_case.id,
            view_type='invalid',
            visible=True
        )
        self.assertEqual(count, 0, "无效视图类型应该返回0")
        
        print("✅ 测试25通过: 批量更新无效视图类型")
    
    def test_26_update_step_view_config_partial(self):
        """测试部分更新步骤视图配置"""
        step = self.db.query(TestStep).filter(TestStep.test_case_id == self.test_case.id).first()
        
        # 只更新业务视图
        result = self.service.update_step_view_config(
            step.id,
            is_business_view=0
        )
        self.assertTrue(result)
        
        self.db.refresh(step)
        self.assertEqual(step.is_business_view, 0)
        
        # 恢复
        self.service.update_step_view_config(step.id, is_business_view=1, is_technical_view=1)
        
        print("✅ 测试26通过: 部分更新步骤视图配置")
    
    def test_27_update_nonexistent_step(self):
        """测试更新不存在的步骤"""
        result = self.service.update_step_view_config(
            999999,
            is_business_view=0
        )
        self.assertFalse(result, "更新不存在的步骤应该失败")
        
        print("✅ 测试27通过: 更新不存在步骤")
    
    def test_28_export_functional_excel_empty(self):
        """测试导出空用例列表到功能用例Excel"""
        import tempfile
        
        temp_dir = tempfile.mkdtemp()
        temp_path = os.path.join(temp_dir, 'empty.xlsx')
        
        try:
            result = self.service.export_to_functional_excel([], temp_path)
            self.assertFalse(result, "空列表应该导出失败")
        finally:
            import shutil
            try:
                shutil.rmtree(temp_dir)
            except:
                pass
        
        print("✅ 测试28通过: 导出空用例列表")
    
    def test_29_validate_excel_file_not_exist(self):
        """测试验证不存在的Excel文件"""
        result = self.service.validate_excel_format('/nonexistent/file.xlsx')
        self.assertFalse(result['valid'])
        self.assertIn('文件不存在', result['errors'])
        
        print("✅ 测试29通过: 验证不存在的文件")
    
    def test_30_validate_functional_excel_invalid(self):
        """测试验证无效的功能用例Excel"""
        import tempfile
        import pandas as pd
        
        temp_dir = tempfile.mkdtemp()
        temp_path = os.path.join(temp_dir, 'invalid_func.xlsx')
        
        try:
            # 创建缺少必需列的Excel
            df = pd.DataFrame([{'col1': 'value1'}])
            df.to_excel(temp_path, index=False)
            
            result = self.service.validate_functional_excel(temp_path)
            self.assertFalse(result['valid'])
            
            print("✅ 测试30通过: 验证无效功能用例Excel")
        finally:
            import shutil
            try:
                shutil.rmtree(temp_dir)
            except:
                pass
    
    def test_31_build_functional_steps_empty(self):
        """测试构建功能用例步骤（空列表）"""
        result = self.service._build_functional_steps([])
        self.assertEqual(result, "", "空列表应该返回空字符串")
        print("✅ 测试31通过: 构建功能用例步骤空列表")
    
    def test_32_build_functional_expected_empty(self):
        """测试构建功能用例预期结果（空列表）"""
        result = self.service._build_functional_expected([])
        self.assertEqual(result, "", "空列表应该返回空字符串")
        print("✅ 测试32通过: 构建功能用例预期结果空列表")
    
    def test_33_export_to_excel_exception(self):
        """测试导出Excel异常处理"""
        # 测试无效的文件路径（目录不存在）
        result = self.service.export_to_excel(self.test_case.id, '/invalid/path/test.xlsx')
        self.assertFalse(result, "无效路径应该返回False")
        print("✅ 测试33通过: 导出Excel异常处理")
    
    def test_34_export_to_functional_excel_exception(self):
        """测试导出功能用例Excel异常处理"""
        result = self.service.export_to_functional_excel([self.test_case.id], '/invalid/path/test.xlsx')
        self.assertFalse(result, "无效路径应该返回False")
        print("✅ 测试34通过: 导出功能用例Excel异常处理")
    
    def test_35_import_from_excel_exception(self):
        """测试导入Excel异常处理"""
        result = self.service.import_from_excel('/nonexistent/file.xlsx', self.test_project.id)
        self.assertIsNone(result, "不存在的文件应该返回None")
        print("✅ 测试35通过: 导入Excel异常处理")
    
    def test_36_import_functional_excel_exception(self):
        """测试导入功能用例Excel异常处理"""
        result = self.service.import_functional_excel('/nonexistent/file.xlsx', self.test_project.id)
        self.assertEqual(result, [], "不存在的文件应该返回空列表")
        print("✅ 测试36通过: 导入功能用例Excel异常处理")
    
    def test_37_import_from_excel_invalid_project(self):
        """测试导入Excel到不存在的项目"""
        import tempfile
        import pandas as pd
        
        temp_dir = tempfile.mkdtemp()
        temp_path = os.path.join(temp_dir, 'test.xlsx')
        
        try:
            with pd.ExcelWriter(temp_path, engine='openpyxl') as writer:
                pd.DataFrame([{'用例编号': 'TC001', '用例标题': '测试'}]).to_excel(writer, sheet_name='用例信息', index=False)
                pd.DataFrame([{'步骤编号': 1, '操作步骤': '操作', '预期结果': '预期'}]).to_excel(writer, sheet_name='测试步骤', index=False)
            
            result = self.service.import_from_excel(temp_path, 999999)
            self.assertIsNone(result, "不存在的项目应该返回None")
            print("✅ 测试37通过: 导入到不存在的项目")
        finally:
            import shutil
            try:
                shutil.rmtree(temp_dir)
            except:
                pass
    
    def test_38_get_technical_view_empty_steps(self):
        """测试获取技术视图（空步骤）"""
        # 创建没有步骤的用例
        empty_case = TestCase(
            project_id=self.test_project.id,
            case_no=f"TC_TECH_EMPTY_{int(datetime.now().timestamp())}",
            module="技术视图空测试",
            title="技术视图空用例",
            precondition="",
            expected_result="",
            priority=2,
            case_type="UI",
            generate_status=1,
            steps_json=[]
        )
        self.db.add(empty_case)
        self.db.commit()
        self.db.refresh(empty_case)
        
        try:
            view = self.service.get_technical_view(empty_case.id)
            self.assertIsNotNone(view)
            self.assertEqual(view['case_id'], empty_case.id)
            self.assertEqual(len(view['steps']), 0)
            self.assertEqual(view['locator_coverage'], 0.0)
            print("✅ 测试38通过: 获取技术视图空步骤")
        finally:
            self.db.query(TestCase).filter(TestCase.id == empty_case.id).delete(synchronize_session=False)
            self.db.commit()
    
    def test_39_get_locator_coverage_with_locators(self):
        """测试获取定位覆盖率（有定位器）"""
        coverage = self.service.get_locator_coverage(self.test_case.id)
        
        self.assertIn('total_steps', coverage)
        self.assertIn('located_steps', coverage)
        self.assertIn('coverage_percentage', coverage)
        self.assertEqual(coverage['total_steps'], 3)
        # 只有第一个步骤有定位器
        self.assertEqual(coverage['located_steps'], 1)
        self.assertAlmostEqual(coverage['coverage_percentage'], 33.33, places=1)
        print("✅ 测试39通过: 获取定位覆盖率有定位器")
    
    def test_40_export_business_view_nonexistent(self):
        """测试导出不存在的用例业务视图"""
        result = self.service.export_business_view_to_markdown(999999)
        self.assertEqual(result, "", "不存在的用例应该返回空字符串")
        print("✅ 测试40通过: 导出不存在的用例业务视图")
    
    def test_41_export_business_html_nonexistent(self):
        """测试导出不存在的用例HTML视图"""
        result = self.service.export_business_view_to_html(999999)
        self.assertEqual(result, "", "不存在的用例应该返回空字符串")
        print("✅ 测试41通过: 导出不存在的用例HTML视图")
    
    def test_42_import_functional_excel_invalid_format(self):
        """测试导入格式错误的功能用例Excel"""
        import tempfile
        import pandas as pd
        
        temp_dir = tempfile.mkdtemp()
        temp_path = os.path.join(temp_dir, 'invalid.xlsx')
        
        try:
            # 创建缺少必需列的Excel
            df = pd.DataFrame([{'col1': 'value1'}])
            df.to_excel(temp_path, index=False)
            
            result = self.service.import_functional_excel(temp_path, self.test_project.id)
            self.assertEqual(result, [], "无效格式应该返回空列表")
            print("✅ 测试42通过: 导入格式错误的功能用例Excel")
        finally:
            import shutil
            try:
                shutil.rmtree(temp_dir)
            except:
                pass
    
    def test_43_generate_unique_case_no_collision(self):
        """测试生成唯一用例编号（处理冲突）"""
        # 使用已存在的用例编号，应该生成新的唯一编号
        existing_no = self.test_case.case_no
        new_no = self.service._generate_unique_case_no(self.test_project.id, existing_no)
        
        # 新编号应该与旧编号不同
        self.assertNotEqual(new_no, existing_no, "应该生成不同的编号")
        # 新编号应该包含项目ID
        self.assertIn(str(self.test_project.id), new_no, "应该包含项目ID")
        print("✅ 测试43通过: 生成唯一用例编号冲突处理")
    
    def test_44_export_technical_view_to_python_with_input(self):
        """测试导出技术视图为Python脚本（包含输入操作）"""
        # 获取第一个已有步骤并更新为输入操作
        step = self.db.query(TestStep).filter(
            TestStep.test_case_id == self.test_case.id
        ).first()
        
        original_action = step.action
        step.action = '输入"测试文本"到搜索框'
        step.has_locator = 1
        step.locator_status = "located"
        self.db.flush()
        
        # 更新定位器
        locator = self.db.query(ElementLocator).filter(
            ElementLocator.step_id == step.id
        ).first()
        if locator:
            locator.css_selector = "#search-input"
        self.db.commit()
        
        try:
            result = self.service.export_technical_view_to_python(self.test_case.id)
            # 检查是否包含fill操作
            self.assertIn('await page.fill', result, "应该包含fill操作")
            self.assertIn('#search-input', result, "应该包含CSS选择器")
            print("✅ 测试44通过: 导出技术视图Python脚本（输入操作）")
        finally:
            step.action = original_action
            if locator:
                locator.css_selector = "#element1"
            self.db.commit()
    
    def test_45_export_technical_view_to_python_without_locator(self):
        """测试导出技术视图为Python脚本（无定位器步骤）"""
        # 创建一个没有定位器的步骤
        no_locator_step = TestStep(
            test_case_id=self.test_case.id,
            step_number=11,
            action="等待页面加载",
            expected_result="页面加载完成",
            is_business_view=1,
            is_technical_view=1,
            has_locator=0,
            locator_status="pending"
        )
        self.db.add(no_locator_step)
        self.db.commit()
        
        try:
            result = self.service.export_technical_view_to_python(self.test_case.id)
            self.assertIn('TODO: 需要添加元素定位', result, "应该包含TODO注释")
            print("✅ 测试45通过: 导出技术视图Python脚本（无定位器）")
        finally:
            self.db.query(TestStep).filter(TestStep.id == no_locator_step.id).delete(synchronize_session=False)
            self.db.commit()
    
    def test_46_get_view_statistics_with_both_views(self):
        """测试获取视图统计（双视图步骤）"""
        # 创建一个只在业务视图显示的步骤
        business_only_step = TestStep(
            test_case_id=self.test_case.id,
            step_number=12,
            action="业务专用步骤",
            expected_result="业务预期",
            is_business_view=1,
            is_technical_view=0,
            has_locator=0,
            locator_status="pending"
        )
        self.db.add(business_only_step)
        self.db.commit()
        
        try:
            stats = self.service.get_view_statistics(self.test_case.id)
            self.assertGreaterEqual(stats['total_steps'], 4, "总步骤数应该>=4")
            print("✅ 测试46通过: 获取视图统计（双视图步骤）")
        finally:
            self.db.query(TestStep).filter(TestStep.id == business_only_step.id).delete(synchronize_session=False)
            self.db.commit()
    
    def test_47_validate_excel_format_missing_required_columns(self):
        """测试验证Excel格式（缺少必需列）"""
        import tempfile
        import pandas as pd
        
        temp_dir = tempfile.mkdtemp()
        temp_path = os.path.join(temp_dir, 'missing_cols.xlsx')
        
        try:
            with pd.ExcelWriter(temp_path, engine='openpyxl') as writer:
                # 用例信息缺少用例标题
                pd.DataFrame([{'用例编号': 'TC001'}]).to_excel(writer, sheet_name='用例信息', index=False)
                # 测试步骤缺少操作步骤
                pd.DataFrame([{'步骤编号': 1, '预期结果': '预期'}]).to_excel(writer, sheet_name='测试步骤', index=False)
            
            result = self.service.validate_excel_format(temp_path)
            self.assertFalse(result['valid'], "缺少必需列应该验证失败")
            self.assertTrue(len(result['errors']) > 0, "应该有错误信息")
            print("✅ 测试47通过: 验证Excel格式缺少必需列")
        finally:
            import shutil
            try:
                shutil.rmtree(temp_dir)
            except:
                pass
    
    def test_48_import_from_excel_with_locator(self):
        """测试从Excel导入（包含定位器信息）"""
        import tempfile
        import pandas as pd
        
        temp_dir = tempfile.mkdtemp()
        temp_path = os.path.join(temp_dir, 'import_with_locator.xlsx')
        
        try:
            with pd.ExcelWriter(temp_path, engine='openpyxl') as writer:
                case_df = pd.DataFrame([{
                    '用例编号': f'LOC_{int(datetime.now().timestamp())}',
                    '用例标题': '带定位器的用例',
                    '所属模块': '定位器测试',
                    '前置条件': '前置',
                    '预期结果': '预期',
                    '优先级': 1,
                    '总步骤数': 1
                }])
                case_df.to_excel(writer, sheet_name='用例信息', index=False)
                
                steps_df = pd.DataFrame([{
                    '步骤编号': 1,
                    '操作步骤': '点击按钮',
                    '预期结果': '按钮被点击',
                    '业务视图': '是',
                    '技术视图': '是',
                    '已定位': '是',
                    '定位状态': 'located',
                    'CSS选择器': '#btn-submit',
                    'XPath': '//button[@id="submit"]',
                    '元素类型': 'button'
                }])
                steps_df.to_excel(writer, sheet_name='测试步骤', index=False)
            
            case_id = self.service.import_from_excel(temp_path, self.test_project.id)
            self.assertIsNotNone(case_id, "导入应该成功")
            
            # 验证定位器是否被创建
            imported_case = self.db.query(TestCase).filter(TestCase.id == case_id).first()
            steps = imported_case.test_steps
            self.assertEqual(len(steps), 1)
            
            # 清理
            self.db.query(TestStep).filter(TestStep.test_case_id == case_id).delete(synchronize_session=False)
            self.db.query(TestCase).filter(TestCase.id == case_id).delete(synchronize_session=False)
            self.db.commit()
            
            print("✅ 测试48通过: 从Excel导入带定位器")
        finally:
            import shutil
            try:
                shutil.rmtree(temp_dir)
            except:
                pass


    def test_49_business_test_case_view_to_dict(self):
        """测试BusinessTestCaseView.to_dict方法"""
        from app.services.test_case_view_service import BusinessTestCaseView, BusinessStepView
        
        # 创建业务视图
        view = BusinessTestCaseView(
            case_id=1,
            case_no="TC001",
            title="测试用例",
            description="描述",
            precondition="前置条件",
            steps=[
                BusinessStepView(step_number=1, action="操作1", expected_result="预期1"),
                BusinessStepView(step_number=2, action="操作2", expected_result="预期2")
            ]
        )
        
        result = view.to_dict()
        
        self.assertEqual(result['case_id'], 1)
        self.assertEqual(result['case_no'], "TC001")
        self.assertEqual(result['title'], "测试用例")
        self.assertEqual(result['description'], "描述")
        self.assertEqual(result['precondition'], "前置条件")
        self.assertEqual(len(result['steps']), 2)
        self.assertEqual(result['total_steps'], 2)
        
        print("✅ 测试49通过: BusinessTestCaseView.to_dict")
    
    def test_50_technical_test_case_view_to_dict(self):
        """测试TechnicalTestCaseView.to_dict方法"""
        from app.services.test_case_view_service import TechnicalTestCaseView, TechnicalStepView
        
        # 创建技术视图
        view = TechnicalTestCaseView(
            case_id=1,
            case_no="TC001",
            title="测试用例",
            steps=[
                TechnicalStepView(step_number=1, action="操作1", expected_result="预期1", has_locator=True, locator_status="located"),
                TechnicalStepView(step_number=2, action="操作2", expected_result="预期2", has_locator=False, locator_status="pending")
            ],
            locator_coverage=50.0
        )
        
        result = view.to_dict()
        
        self.assertEqual(result['case_id'], 1)
        self.assertEqual(result['case_no'], "TC001")
        self.assertEqual(result['title'], "测试用例")
        self.assertEqual(len(result['steps']), 2)
        self.assertEqual(result['total_steps'], 2)
        self.assertEqual(result['locator_coverage'], "50.0%")
        
        print("✅ 测试50通过: TechnicalTestCaseView.to_dict")


if __name__ == '__main__':
    unittest.main(verbosity=2)

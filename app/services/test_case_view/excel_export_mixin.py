"""Excel导出Mixin - 标准格式和功能用例格式的Excel导出。"""
from typing import List
from loguru import logger

from app.models.test_case import TestCase, TestStep
from app.models.element_locator import ElementLocator


class ExcelExportMixin:
    """Excel导出：标准双Sheet格式、第三方功能用例格式。"""

    def export_to_excel(self, test_case_id: int, file_path: str) -> bool:
        """导出单条用例为标准双 Sheet Excel（用例信息 + 测试步骤），纯 openpyxl 实现。"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment

            test_case = self.db.query(TestCase).filter(TestCase.id == test_case_id, TestCase.is_deleted.is_(False)).first()
            if not test_case:
                logger.warning(f"测试用例不存在: {test_case_id}")
                return False

            steps = self.db.query(TestStep).filter(
                TestStep.test_case_id == test_case_id
            ).order_by(TestStep.step_number).all()

            step_ids = [s.id for s in steps]
            locators = {}
            if step_ids:
                locator_list = self.db.query(ElementLocator).filter(
                    ElementLocator.step_id.in_(step_ids)
                ).all()
                locators = {l.step_id: l for l in locator_list}

            wb = Workbook()

            # ====== Sheet 1: 用例信息 ======
            ws_info = wb.active
            ws_info.title = "用例信息"
            info_columns = [
                ("用例编号", test_case.case_no or ""),
                ("用例标题", test_case.title or ""),
                ("所属模块", test_case.module or ""),
                ("前置条件", test_case.precondition or ""),
                ("预期结果", test_case.expected_result or ""),
                ("优先级", test_case.priority),
                ("总步骤数", len(steps)),
            ]
            header_font = Font(bold=True)
            header_fill = PatternFill(start_color="D9D2E9", end_color="D9D2E9", fill_type="solid")
            for col_idx, (header, _) in enumerate(info_columns, 1):
                cell = ws_info.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
            for col_idx, (_, value) in enumerate(info_columns, 1):
                ws_info.cell(row=2, column=col_idx, value=value)
            ws_info.column_dimensions["A"].width = 14
            ws_info.column_dimensions["B"].width = 40
            ws_info.column_dimensions["C"].width = 18
            ws_info.column_dimensions["D"].width = 30
            ws_info.column_dimensions["E"].width = 30

            # ====== Sheet 2: 测试步骤 ======
            ws_steps = wb.create_sheet(title="测试步骤")
            step_headers = [
                "步骤编号", "操作步骤", "预期结果", "业务视图", "技术视图",
                "已定位", "定位状态", "CSS选择器", "XPath", "元素类型",
            ]
            for col_idx, header in enumerate(step_headers, 1):
                cell = ws_steps.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
            wrap = Alignment(wrap_text=True, vertical="top")
            for row_idx, step in enumerate(steps, 2):
                locator = locators.get(step.id)
                row_values = [
                    step.step_number,
                    step.action or "",
                    step.expected_result or "",
                    "是" if step.is_business_view == 1 else "否",
                    "是" if step.is_technical_view == 1 else "否",
                    "是" if step.has_locator == 1 else "否",
                    step.locator_status or "",
                    locator.css_selector if locator else "",
                    locator.xpath if locator else "",
                    locator.element_type if locator else "",
                ]
                for col_idx, value in enumerate(row_values, 1):
                    cell = ws_steps.cell(row=row_idx, column=col_idx, value=value)
                    if col_idx in (2, 3, 8, 9):
                        cell.alignment = wrap
            widths = [10, 30, 30, 10, 10, 10, 12, 24, 24, 14]
            for col_idx, w in enumerate(widths, 1):
                ws_steps.column_dimensions[ws_steps.cell(row=1, column=col_idx).column_letter].width = w

            wb.save(file_path)
            logger.info(f"测试用例 {test_case_id} 导出Excel成功: {file_path}")
            return True
        except Exception as e:
            logger.error(f"导出Excel失败: {e}", exc_info=True)
            return False

    def export_to_functional_excel(self, test_case_ids: List[int], file_path: str) -> bool:
        try:
            from collections import OrderedDict
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

            if not test_case_ids:
                logger.warning("未提供测试用例ID列表")
                return False

            test_cases = self.db.query(TestCase).filter(
                TestCase.id.in_(test_case_ids),
                TestCase.is_deleted.is_(False)
            ).all()

            if not test_cases:
                logger.warning(f"未找到测试用例: {test_case_ids}")
                return False

            # 按调用方传入的 ID 顺序排列，保证导出顺序稳定
            id_order = {tid: idx for idx, tid in enumerate(test_case_ids)}
            test_cases.sort(key=lambda tc: id_order.get(tc.id, 0))

            # 按模块分组，保持原始顺序
            modules: OrderedDict = OrderedDict()
            for tc in test_cases:
                module = tc.module or "默认模块"
                if module not in modules:
                    modules[module] = []
                modules[module].append(tc)

            wb = Workbook()
            ws = wb.active
            ws.title = "测试用例"

            # 样式定义
            header_font = Font(bold=True, size=11)
            header_fill = PatternFill(start_color="D9D2E9", end_color="D9D2E9", fill_type="solid")
            module_font = Font(bold=True, size=11)
            module_fill = PatternFill(start_color="E8E0F0", end_color="E8E0F0", fill_type="solid")
            thin_border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )
            wrap_alignment = Alignment(wrap_text=True, vertical="top")
            center_alignment = Alignment(horizontal="center", vertical="top")

            # 列定义
            columns = ["用例序号", "优先级", "用例描述", "初始条件", "操作步骤", "期望结果", "测试结果", "测试人"]
            col_widths = [12, 8, 30, 25, 35, 35, 12, 10]

            # 写表头
            for col_idx, col_name in enumerate(columns, 1):
                cell = ws.cell(row=1, column=col_idx, value=col_name)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for col_idx, width in enumerate(col_widths, 1):
                ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

            current_row = 2
            priority_map = {1: "P0", 2: "P2", 3: "P3"}

            # 批量查询所有步骤，避免 N+1
            all_case_ids = [tc.id for tc in test_cases]
            all_steps = self.db.query(TestStep).filter(
                TestStep.test_case_id.in_(all_case_ids)
            ).order_by(TestStep.test_case_id, TestStep.step_number).all()
            steps_by_case: dict = {}
            for s in all_steps:
                steps_by_case.setdefault(s.test_case_id, []).append(s)

            for module_name, cases in modules.items():
                # 写模块分组标题行
                ws.merge_cells(
                    start_row=current_row, start_column=1,
                    end_row=current_row, end_column=len(columns),
                )
                module_cell = ws.cell(row=current_row, column=1, value=module_name)
                module_cell.font = module_font
                module_cell.fill = module_fill
                module_cell.alignment = Alignment(horizontal="center", vertical="center")
                for ci in range(1, len(columns) + 1):
                    ws.cell(row=current_row, column=ci).border = thin_border
                current_row += 1

                for seq, tc in enumerate(cases, 1):
                    steps = steps_by_case.get(tc.id, [])

                    step_desc = self._build_functional_steps(steps)
                    expected = self._build_functional_expected(steps)
                    if not expected:
                        expected = tc.expected_result or ""
                    priority = priority_map.get(tc.priority, "P2")

                    row_data = [
                        seq,
                        priority,
                        tc.title or "",
                        tc.precondition or "",
                        step_desc,
                        expected,
                        "",
                        "",
                    ]
                    for col_idx, value in enumerate(row_data, 1):
                        cell = ws.cell(row=current_row, column=col_idx, value=value)
                        cell.border = thin_border
                        if col_idx in (1, 2):
                            cell.alignment = center_alignment
                        else:
                            cell.alignment = wrap_alignment
                    current_row += 1

            wb.save(file_path)
            total = sum(len(v) for v in modules.values())
            logger.info(f"导出功能用例Excel成功: {file_path}, 共{total}条用例")
            return True
        except Exception as e:
            logger.error(f"导出功能用例Excel失败: {e}")
            return False

    def _build_functional_steps(self, steps: List[TestStep]) -> str:
        if not steps:
            return ""
        return "\n".join(f"[{s.step_number}] {s.action}" for s in steps)

    def _build_functional_expected(self, steps: List[TestStep]) -> str:
        if not steps:
            return ""
        items = []
        for s in steps:
            if s.expected_result:
                items.append(f"[{s.step_number}] {s.expected_result}")
        return "\n".join(items)

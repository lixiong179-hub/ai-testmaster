"""文件内容提取 - 文本与Office文档提取器
"""
import os
import re
from typing import Optional
from loguru import logger

from app.models.project import ProjectFile


async def extract_text_file(file: ProjectFile) -> Optional[str]:
    """提取文本文件内容"""
    if not os.path.exists(file.file_url):
        logger.warning(f"文件不存在: {file.file_url}")
        return None
    try:
        with open(file.file_url, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.read()
        return clean_text(content) if len(content) > 0 else None
    except Exception as e:
        logger.error(f"读取文本文件失败: {e}")
        return None


async def extract_docx(file: ProjectFile) -> Optional[str]:
    """提取 Word 文档内容"""
    try:
        from docx import Document
        if not os.path.exists(file.file_url):
            return None
        doc = Document(file.file_url)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        content = '\n'.join(paragraphs)
        return clean_text(content) if len(content) > 0 else None
    except ImportError:
        logger.warning("python-docx 未安装，无法解析 Word 文档")
        return "【Word文档】请安装 python-docx 库以提取文本内容"
    except Exception as e:
        logger.error(f"解析 Word 文档失败: {e}")
        return None


async def extract_pdf(file: ProjectFile) -> Optional[str]:
    """提取 PDF 文档内容"""
    try:
        import pypdf
        if not os.path.exists(file.file_url):
            return None
        reader = pypdf.PdfReader(file.file_url)
        text_parts = []
        for page_num, page in enumerate(reader.pages, 1):
            text = page.extract_text()
            if text:
                text_parts.append(f"--- 第{page_num}页 ---\n{text}")
        content = '\n'.join(text_parts)
        return clean_text(content) if len(content) > 0 else None
    except ImportError:
        logger.warning("pypdf 未安装，无法解析 PDF 文档")
        return "【PDF文档】请安装 pypdf 库以提取文本内容"
    except Exception as e:
        logger.error(f"解析 PDF 文档失败: {e}")
        return None


async def extract_excel(file: ProjectFile) -> Optional[str]:
    """提取 Excel 表格内容"""
    try:
        import openpyxl
        if not os.path.exists(file.file_url):
            return None
        wb = openpyxl.load_workbook(file.file_url, data_only=True)
        text_parts = []
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            text_parts.append(f"=== 工作表: {sheet_name} ===")
            for row in sheet.iter_rows(values_only=True):
                row_data = [str(cell) if cell is not None else '' for cell in row]
                if any(cell.strip() for cell in row_data):
                    text_parts.append(' | '.join(row_data))
        content = '\n'.join(text_parts)
        return clean_text(content) if len(content) > 0 else None
    except ImportError:
        logger.warning("openpyxl 未安装，无法解析 Excel 文档")
        return "【Excel文档】请安装 openpyxl 库以提取表格内容"
    except Exception as e:
        logger.error(f"解析 Excel 文档失败: {e}")
        return None


def clean_text(text: str, max_length: int = 50000) -> str:
    """清理文本内容"""
    if not text:
        return ""
    text = re.sub(r'\r\n|\r|\n', '\n', text)
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    if len(text) > max_length:
        text = text[:max_length] + "\n\n...（内容过长已截断）"
    return text.strip()
